"""Persistent DoubleTick chat import jobs used by the FastAPI application."""

from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import closing, contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

from chat_search_index import (
    DB_PATH,
    DISCLAIMER,
    contains_stop_phrase,
    refresh_analysis_artifacts_preserving_enrichment,
)
from chat_search_taxonomy import normalize_text


DATA = Path(os.environ.get("APP_DATA_DIR", "data"))
DATA.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR = DATA / "import_uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
CONTROL_DB_PATH = DATA / "chat_imports.sqlite3"
DOUBLETICK_URL = "https://public.doubletick.io/chat-messages"
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_WORKERS = 5
_RUNNING_JOBS: dict[str, threading.Event] = {}
_RUNNING_LOCK = threading.Lock()
_HTTP_LOCAL = threading.local()


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


@contextmanager
def _connection():
    conn = sqlite3.connect(CONTROL_DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA busy_timeout = 30000")
    try:
        with conn:
            yield conn
    finally:
        conn.close()


def initialize_import_store() -> None:
    with _connection() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS import_uploads (
                upload_id TEXT PRIMARY KEY,
                original_filename TEXT NOT NULL,
                saved_path TEXT NOT NULL,
                file_bytes INTEGER NOT NULL,
                workbook_json TEXT NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS import_jobs (
                job_id TEXT PRIMARY KEY,
                upload_id TEXT NOT NULL,
                sheet_name TEXT NOT NULL,
                waba_column TEXT NOT NULL,
                customer_column TEXT NOT NULL,
                start_date TEXT,
                end_date TEXT,
                workers INTEGER NOT NULL,
                request_delay_ms INTEGER NOT NULL,
                status TEXT NOT NULL,
                total_pairs INTEGER NOT NULL DEFAULT 0,
                completed_pairs INTEGER NOT NULL DEFAULT 0,
                successful_pairs INTEGER NOT NULL DEFAULT 0,
                failed_pairs INTEGER NOT NULL DEFAULT 0,
                fetched_messages INTEGER NOT NULL DEFAULT 0,
                inserted_messages INTEGER NOT NULL DEFAULT 0,
                duplicate_messages INTEGER NOT NULL DEFAULT 0,
                skipped_messages INTEGER NOT NULL DEFAULT 0,
                last_error TEXT,
                created_at TEXT NOT NULL,
                started_at TEXT,
                finished_at TEXT,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(upload_id) REFERENCES import_uploads(upload_id)
            );

            CREATE TABLE IF NOT EXISTS import_pairs (
                job_id TEXT NOT NULL,
                row_number INTEGER NOT NULL,
                waba_number TEXT NOT NULL,
                customer_number TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending',
                attempts INTEGER NOT NULL DEFAULT 0,
                message_count INTEGER NOT NULL DEFAULT 0,
                error TEXT,
                PRIMARY KEY(job_id, row_number)
            );

            CREATE TABLE IF NOT EXISTS import_messages (
                job_id TEXT NOT NULL,
                fingerprint TEXT NOT NULL,
                conversation_id TEXT NOT NULL,
                message_timestamp INTEGER NOT NULL,
                message_datetime TEXT NOT NULL,
                he_number TEXT NOT NULL,
                customer_number TEXT NOT NULL,
                sender_number TEXT NOT NULL,
                sender_type TEXT NOT NULL,
                message_type TEXT NOT NULL,
                message_content TEXT NOT NULL,
                message_content_lower TEXT NOT NULL,
                message_content_normalized TEXT NOT NULL,
                source_file TEXT NOT NULL,
                PRIMARY KEY(job_id, fingerprint)
            );

            CREATE INDEX IF NOT EXISTS idx_import_pairs_status
            ON import_pairs(job_id, status, row_number);
            CREATE INDEX IF NOT EXISTS idx_import_messages_job
            ON import_messages(job_id, conversation_id);
            """
        )
        conn.execute(
            """
            UPDATE import_jobs
            SET status = 'interrupted',
                last_error = 'Server restarted. Re-enter the authorization key and resume.',
                updated_at = ?
            WHERE status IN ('queued', 'running', 'cancel_requested', 'merging', 'indexing')
            """,
            (_now(),),
        )
        conn.execute(
            "UPDATE import_pairs SET status = 'pending' WHERE status = 'in_progress'"
        )


def _safe_filename(filename: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name)
    return cleaned[:120] or "customers.xlsx"


def _workbook_info(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = [str(value).strip() if value is not None else "" for value in header_row]
            nonempty_headers = [value for value in headers if value]
            sheets.append(
                {
                    "name": worksheet.title,
                    "headers": nonempty_headers,
                    "row_count": max(int(worksheet.max_row or 1) - 1, 0),
                    "suggested_waba_column": _find_header(
                        nonempty_headers, ["he phone", "waba number", "waba", "business number"]
                    ),
                    "suggested_customer_column": _find_header(
                        nonempty_headers, ["phone number", "customer number", "customer phone", "mobile"]
                    ),
                }
            )
    finally:
        workbook.close()
    if not sheets:
        raise ValueError("The workbook has no readable sheets")
    return {"sheets": sheets}


def _find_header(headers: list[str], candidates: list[str]) -> str | None:
    lookup = {header.casefold(): header for header in headers}
    for candidate in candidates:
        if candidate in lookup:
            return lookup[candidate]
    return None


def save_upload(filename: str, content: bytes) -> dict[str, Any]:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Upload an .xlsx workbook")
    if not content:
        raise ValueError("The uploaded workbook is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("Workbook exceeds the 25 MB upload limit")

    upload_id = uuid.uuid4().hex
    path = UPLOAD_DIR / f"{upload_id}_{_safe_filename(filename)}"
    path.write_bytes(content)
    try:
        workbook_info = _workbook_info(path)
    except Exception:
        path.unlink(missing_ok=True)
        raise

    created_at = _now()
    with _connection() as conn:
        conn.execute(
            """
            INSERT INTO import_uploads(
                upload_id, original_filename, saved_path, file_bytes, workbook_json, created_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (upload_id, filename, str(path), len(content), json.dumps(workbook_info), created_at),
        )
    return {
        "upload_id": upload_id,
        "filename": filename,
        "file_bytes": len(content),
        "created_at": created_at,
        **workbook_info,
    }


def _normalize_phone(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = re.sub(r"\D", "", str(value).strip())
    if not digits:
        return None
    if len(digits) == 10:
        digits = "91" + digits
    if len(digits) < 8 or len(digits) > 15:
        return None
    return "+" + digits


def _load_pairs(
    upload_path: Path,
    sheet_name: str,
    waba_column: str,
    customer_column: str,
) -> list[tuple[int, str, str]]:
    workbook = load_workbook(upload_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("The selected sheet is empty")
        header_lookup = {
            str(value).strip(): index for index, value in enumerate(headers) if value is not None
        }
        if waba_column not in header_lookup or customer_column not in header_lookup:
            raise ValueError("The selected phone columns are not present in the sheet")
        waba_index = header_lookup[waba_column]
        customer_index = header_lookup[customer_column]
        pairs: list[tuple[int, str, str]] = []
        seen: set[tuple[str, str]] = set()
        for row_number, row in enumerate(rows, start=2):
            waba = _normalize_phone(row[waba_index] if waba_index < len(row) else None)
            customer = _normalize_phone(row[customer_index] if customer_index < len(row) else None)
            if not waba or not customer or (waba, customer) in seen:
                continue
            seen.add((waba, customer))
            pairs.append((row_number, waba, customer))
        return pairs
    finally:
        workbook.close()


def _validate_iso_date(value: Any, field: str) -> str | None:
    if value in (None, ""):
        return None
    try:
        parsed = datetime.strptime(str(value), "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError(f"{field} must use YYYY-MM-DD") from exc
    return parsed.strftime("%d-%m-%Y")


def create_job(payload: dict[str, Any], auth_key: str) -> dict[str, Any]:
    if not auth_key.strip():
        raise ValueError("DoubleTick authorization key is required")
    upload_id = str(payload.get("upload_id") or "").strip()
    sheet_name = str(payload.get("sheet_name") or "").strip()
    waba_column = str(payload.get("waba_column") or "").strip()
    customer_column = str(payload.get("customer_column") or "").strip()
    if not all([upload_id, sheet_name, waba_column, customer_column]):
        raise ValueError("Upload, sheet, WABA column, and customer column are required")

    workers = max(1, min(int(payload.get("workers") or 3), MAX_WORKERS))
    request_delay_ms = max(0, min(int(payload.get("request_delay_ms") or 100), 5000))
    start_date = _validate_iso_date(payload.get("start_date"), "start_date")
    end_date = _validate_iso_date(payload.get("end_date"), "end_date")
    if start_date and end_date:
        parsed_start = datetime.strptime(start_date, "%d-%m-%Y")
        parsed_end = datetime.strptime(end_date, "%d-%m-%Y")
        if parsed_start > parsed_end:
            raise ValueError("start_date cannot be after end_date")

    with _connection() as conn:
        active = conn.execute(
            """
            SELECT job_id FROM import_jobs
            WHERE status IN ('queued', 'running', 'cancel_requested', 'merging', 'indexing')
            LIMIT 1
            """
        ).fetchone()
        if active:
            raise ValueError("Another chat import is active. Let it finish or cancel it first.")
        upload = conn.execute(
            "SELECT * FROM import_uploads WHERE upload_id = ?", (upload_id,)
        ).fetchone()
    if upload is None:
        raise ValueError("Uploaded workbook was not found")

    pairs = _load_pairs(Path(upload["saved_path"]), sheet_name, waba_column, customer_column)
    if not pairs:
        raise ValueError("No valid WABA/customer phone pairs were found")

    job_id = uuid.uuid4().hex
    created_at = _now()
    with _connection() as conn:
        conn.execute(
            """
            INSERT INTO import_jobs(
                job_id, upload_id, sheet_name, waba_column, customer_column,
                start_date, end_date, workers, request_delay_ms, status,
                total_pairs, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'queued', ?, ?, ?)
            """,
            (
                job_id, upload_id, sheet_name, waba_column, customer_column,
                start_date, end_date, workers, request_delay_ms, len(pairs),
                created_at, created_at,
            ),
        )
        conn.executemany(
            """
            INSERT INTO import_pairs(job_id, row_number, waba_number, customer_number)
            VALUES (?, ?, ?, ?)
            """,
            ((job_id, row_number, waba, customer) for row_number, waba, customer in pairs),
        )
    _launch_job(job_id, auth_key.strip())
    return get_job(job_id)


def _job_payload(row: sqlite3.Row) -> dict[str, Any]:
    result = dict(row)
    total = int(result.get("total_pairs") or 0)
    completed = int(result.get("completed_pairs") or 0)
    result["progress_percent"] = round((completed / total) * 100, 1) if total else 0.0
    result["can_cancel"] = result["status"] in {"queued", "running"}
    result["can_resume"] = result["status"] in {
        "interrupted", "failed", "cancelled", "completed_with_errors"
    }
    return result


def get_job(job_id: str) -> dict[str, Any]:
    with _connection() as conn:
        row = conn.execute(
            """
            SELECT jobs.*, uploads.original_filename,
                (
                    SELECT pairs.error FROM import_pairs pairs
                    WHERE pairs.job_id = jobs.job_id AND pairs.status = 'failed'
                    ORDER BY pairs.row_number LIMIT 1
                ) AS sample_error
            FROM import_jobs jobs
            JOIN import_uploads uploads ON uploads.upload_id = jobs.upload_id
            WHERE jobs.job_id = ?
            """,
            (job_id,),
        ).fetchone()
        if row is None:
            raise KeyError(job_id)
        errors = conn.execute(
            """
            SELECT row_number, waba_number, customer_number, attempts, error
            FROM import_pairs
            WHERE job_id = ? AND status = 'failed'
            ORDER BY row_number
            LIMIT 20
            """,
            (job_id,),
        ).fetchall()
    result = _job_payload(row)
    result["recent_errors"] = [dict(error) for error in errors]
    return result


def list_jobs(limit: int = 20) -> list[dict[str, Any]]:
    with _connection() as conn:
        rows = conn.execute(
            """
            SELECT jobs.*, uploads.original_filename,
                (
                    SELECT pairs.error FROM import_pairs pairs
                    WHERE pairs.job_id = jobs.job_id AND pairs.status = 'failed'
                    ORDER BY pairs.row_number LIMIT 1
                ) AS sample_error
            FROM import_jobs jobs
            JOIN import_uploads uploads ON uploads.upload_id = jobs.upload_id
            ORDER BY jobs.created_at DESC
            LIMIT ?
            """,
            (max(1, min(limit, 100)),),
        ).fetchall()
    return [_job_payload(row) for row in rows]


def cancel_job(job_id: str) -> dict[str, Any]:
    with _RUNNING_LOCK:
        event = _RUNNING_JOBS.get(job_id)
        if event:
            event.set()
    with _connection() as conn:
        row = conn.execute("SELECT status FROM import_jobs WHERE job_id = ?", (job_id,)).fetchone()
        if row is None:
            raise KeyError(job_id)
        if row["status"] in {"queued", "running"}:
            conn.execute(
                "UPDATE import_jobs SET status = 'cancel_requested', updated_at = ? WHERE job_id = ?",
                (_now(), job_id),
            )
    return get_job(job_id)


def resume_job(job_id: str, auth_key: str, retry_failed: bool = True) -> dict[str, Any]:
    if not auth_key.strip():
        raise ValueError("DoubleTick authorization key is required")
    with _RUNNING_LOCK:
        if job_id in _RUNNING_JOBS:
            raise ValueError("This import job is already running")
    with _connection() as conn:
        row = conn.execute("SELECT status FROM import_jobs WHERE job_id = ?", (job_id,)).fetchone()
        if row is None:
            raise KeyError(job_id)
        if row["status"] not in {"interrupted", "failed", "cancelled", "completed_with_errors"}:
            raise ValueError(f"A {row['status']} job cannot be resumed")
        active = conn.execute(
            """
            SELECT job_id FROM import_jobs
            WHERE job_id <> ?
              AND status IN ('queued', 'running', 'cancel_requested', 'merging', 'indexing')
            LIMIT 1
            """,
            (job_id,),
        ).fetchone()
        if active:
            raise ValueError("Another chat import is active. Let it finish or cancel it first.")
        conn.execute(
            "UPDATE import_pairs SET status = 'pending', error = NULL WHERE job_id = ? AND status = 'in_progress'",
            (job_id,),
        )
        if retry_failed:
            conn.execute(
                "UPDATE import_pairs SET status = 'pending', error = NULL WHERE job_id = ? AND status = 'failed'",
                (job_id,),
            )
        conn.execute(
            """
            UPDATE import_jobs
            SET status = 'queued', last_error = NULL, finished_at = NULL, updated_at = ?
            WHERE job_id = ?
            """,
            (_now(), job_id),
        )
    _launch_job(job_id, auth_key.strip())
    return get_job(job_id)


def _launch_job(job_id: str, auth_key: str) -> None:
    with _RUNNING_LOCK:
        if job_id in _RUNNING_JOBS:
            raise ValueError("This import job is already running")
        cancel_event = threading.Event()
        _RUNNING_JOBS[job_id] = cancel_event
    thread = threading.Thread(
        target=_run_job,
        args=(job_id, auth_key, cancel_event),
        name=f"doubletick-import-{job_id[:8]}",
        daemon=True,
    )
    thread.start()


def _extract_message_rows(
    job_id: str,
    waba: str,
    customer: str,
    messages: list[dict[str, Any]],
) -> tuple[list[tuple[Any, ...]], int]:
    stored_he = "'" + waba
    stored_customer = "'" + customer
    conversation_id = f"{stored_he}_{stored_customer}"
    rows: list[tuple[Any, ...]] = []
    skipped = 0
    for item in messages:
        if not isinstance(item, dict):
            skipped += 1
            continue
        message = item.get("message") if isinstance(item.get("message"), dict) else item
        message_type = str(message.get("messageType") or item.get("messageType") or "text").lower()
        text_value = message.get("text") or item.get("text") or ""
        file_url = message.get("fileUrl") or item.get("fileUrl") or ""
        content = "\n\n".join(
            str(value).strip() for value in (text_value, file_url)
            if value is not None and str(value).strip()
        )
        if not content or message_type in {"system", "template"}:
            skipped += 1
            continue
        lowered = content.lower()
        if DISCLAIMER in lowered or contains_stop_phrase(content):
            skipped += 1
            continue
        raw_timestamp = item.get("messageTime") or item.get("timestamp") or item.get("createdAt")
        try:
            timestamp = int(float(raw_timestamp))
            if timestamp > 10_000_000_000:
                timestamp //= 1000
        except (TypeError, ValueError):
            skipped += 1
            continue
        origin = str(item.get("messageOriginType") or item.get("originType") or "").upper()
        sender_type = "HE" if origin in {"USER", "AGENT", "BUSINESS", "WABA"} else "Customer"
        sender_number = stored_he if sender_type == "HE" else stored_customer
        message_datetime = datetime.fromtimestamp(timestamp, tz=timezone.utc).strftime(
            "%Y-%m-%d %H:%M:%S UTC"
        )
        fingerprint_source = "\x1f".join(
            [conversation_id, str(timestamp), sender_type, message_type, content]
        )
        fingerprint = hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest()
        rows.append(
            (
                job_id, fingerprint, conversation_id, timestamp, message_datetime,
                stored_he, stored_customer, sender_number, sender_type, message_type,
                content, lowered, normalize_text(content), f"doubletick:{job_id}",
            )
        )
    return rows, skipped


def _fetch_pair(
    job: dict[str, Any], pair: sqlite3.Row, auth_key: str
) -> tuple[int, list[tuple[Any, ...]], int, str | None, int]:
    params = {
        "wabaNumber": pair["waba_number"].lstrip("+"),
        "customerNumber": pair["customer_number"].lstrip("+"),
    }
    if job.get("start_date"):
        params["startDate"] = job["start_date"]
    if job.get("end_date"):
        params["endDate"] = job["end_date"]
    headers = {"Authorization": auth_key, "Accept": "application/json"}
    last_error: str | None = None
    attempts = 0
    for attempts in range(1, 4):
        try:
            session = getattr(_HTTP_LOCAL, "session", None)
            if session is None:
                session = requests.Session()
                _HTTP_LOCAL.session = session
            response = session.get(
                DOUBLETICK_URL, params=params, headers=headers, timeout=(10, 45)
            )
            if response.status_code == 200:
                payload = response.json()
                messages = payload.get("messages", []) if isinstance(payload, dict) else []
                if not isinstance(messages, list):
                    return pair["row_number"], [], 0, "Response did not contain a messages list", attempts
                rows, skipped = _extract_message_rows(
                    job["job_id"], pair["waba_number"], pair["customer_number"], messages
                )
                return pair["row_number"], rows, skipped, None, attempts
            detail = response.text[:300].replace("\n", " ")
            last_error = f"HTTP {response.status_code}: {detail}"
            if response.status_code not in {408, 429, 500, 502, 503, 504}:
                break
            retry_after = response.headers.get("Retry-After")
            wait_seconds = float(retry_after) if retry_after and retry_after.isdigit() else 2 ** attempts
            time.sleep(min(wait_seconds, 30))
        except (requests.RequestException, ValueError) as exc:
            last_error = f"{type(exc).__name__}: {str(exc)[:300]}"
            if attempts < 3:
                time.sleep(2 ** attempts)
    return pair["row_number"], [], 0, last_error or "Unknown request failure", attempts


def _update_job_counts(conn: sqlite3.Connection, job_id: str) -> None:
    counts = conn.execute(
        """
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN status IN ('succeeded', 'failed') THEN 1 ELSE 0 END) AS completed,
            SUM(CASE WHEN status = 'succeeded' THEN 1 ELSE 0 END) AS successful,
            SUM(CASE WHEN status = 'failed' THEN 1 ELSE 0 END) AS failed,
            SUM(message_count) AS fetched
        FROM import_pairs WHERE job_id = ?
        """,
        (job_id,),
    ).fetchone()
    conn.execute(
        """
        UPDATE import_jobs
        SET total_pairs = ?, completed_pairs = ?, successful_pairs = ?,
            failed_pairs = ?, fetched_messages = ?, updated_at = ?
        WHERE job_id = ?
        """,
        (
            counts["total"] or 0, counts["completed"] or 0, counts["successful"] or 0,
            counts["failed"] or 0, counts["fetched"] or 0, _now(), job_id,
        ),
    )


def _run_job(job_id: str, auth_key: str, cancel_event: threading.Event) -> None:
    try:
        with _connection() as conn:
            job_row = conn.execute("SELECT * FROM import_jobs WHERE job_id = ?", (job_id,)).fetchone()
            if job_row is None:
                return
            job = dict(job_row)
            conn.execute(
                """
                UPDATE import_jobs SET status = 'running', started_at = COALESCE(started_at, ?),
                    updated_at = ? WHERE job_id = ?
                """,
                (_now(), _now(), job_id),
            )

        while not cancel_event.is_set():
            with _connection() as conn:
                pairs = conn.execute(
                    """
                    SELECT * FROM import_pairs
                    WHERE job_id = ? AND status = 'pending'
                    ORDER BY row_number LIMIT 100
                    """,
                    (job_id,),
                ).fetchall()
                if not pairs:
                    break
                conn.executemany(
                    "UPDATE import_pairs SET status = 'in_progress' WHERE job_id = ? AND row_number = ?",
                    ((job_id, pair["row_number"]) for pair in pairs),
                )

            with ThreadPoolExecutor(max_workers=job["workers"]) as executor:
                futures = []
                for pair in pairs:
                    if cancel_event.is_set():
                        break
                    futures.append(executor.submit(_fetch_pair, job, pair, auth_key))
                    if job["request_delay_ms"]:
                        time.sleep(job["request_delay_ms"] / 1000)
                for future in as_completed(futures):
                    row_number, rows, skipped, error, attempts = future.result()
                    with _connection() as conn:
                        if rows:
                            conn.executemany(
                                """
                                INSERT OR IGNORE INTO import_messages(
                                    job_id, fingerprint, conversation_id, message_timestamp,
                                    message_datetime, he_number, customer_number, sender_number,
                                    sender_type, message_type, message_content, message_content_lower,
                                    message_content_normalized, source_file
                                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                """,
                                rows,
                            )
                        conn.execute(
                            """
                            UPDATE import_pairs
                            SET status = ?, attempts = attempts + ?, message_count = ?, error = ?
                            WHERE job_id = ? AND row_number = ?
                            """,
                            (
                                "failed" if error else "succeeded", attempts, len(rows), error,
                                job_id, row_number,
                            ),
                        )
                        conn.execute(
                            "UPDATE import_jobs SET skipped_messages = skipped_messages + ? WHERE job_id = ?",
                            (skipped, job_id),
                        )
                        _update_job_counts(conn, job_id)

        if cancel_event.is_set():
            with _connection() as conn:
                conn.execute(
                    "UPDATE import_pairs SET status = 'pending' WHERE job_id = ? AND status = 'in_progress'",
                    (job_id,),
                )
                conn.execute(
                    "UPDATE import_jobs SET status = 'cancelled', finished_at = ?, updated_at = ? WHERE job_id = ?",
                    (_now(), _now(), job_id),
                )
            return

        _merge_job(job_id)
    except Exception as exc:
        with _connection() as conn:
            conn.execute(
                """
                UPDATE import_jobs SET status = 'failed', last_error = ?, finished_at = ?, updated_at = ?
                WHERE job_id = ?
                """,
                (f"{type(exc).__name__}: {str(exc)[:500]}", _now(), _now(), job_id),
            )
    finally:
        auth_key = ""
        with _RUNNING_LOCK:
            _RUNNING_JOBS.pop(job_id, None)


def _merge_job(job_id: str) -> None:
    with _connection() as conn:
        conn.execute(
            "UPDATE import_jobs SET status = 'merging', updated_at = ? WHERE job_id = ?",
            (_now(), job_id),
        )
        staged = conn.execute(
            "SELECT * FROM import_messages WHERE job_id = ? ORDER BY conversation_id, message_timestamp",
            (job_id,),
        ).fetchall()

    inserted = 0
    duplicates = 0
    with closing(sqlite3.connect(DB_PATH, timeout=120)) as analysis:
        analysis.row_factory = sqlite3.Row
        analysis.execute("PRAGMA busy_timeout = 120000")
        existing_by_conversation: dict[str, set[tuple[Any, ...]]] = {}
        insert_sql = """
            INSERT INTO messages(
                conversation_id, message_timestamp, message_datetime, he_number,
                customer_number, sender_number, sender_type, message_type,
                message_content, message_content_lower, message_content_normalized, source_file
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """
        for row in staged:
            conversation_id = row["conversation_id"]
            if conversation_id not in existing_by_conversation:
                existing_rows = analysis.execute(
                    """
                    SELECT message_timestamp, sender_type, message_type, message_content
                    FROM messages WHERE conversation_id = ?
                    """,
                    (conversation_id,),
                ).fetchall()
                existing_by_conversation[conversation_id] = {
                    (
                        item["message_timestamp"], item["sender_type"],
                        item["message_type"], item["message_content"],
                    )
                    for item in existing_rows
                }
            key = (
                row["message_timestamp"], row["sender_type"], row["message_type"], row["message_content"]
            )
            if key in existing_by_conversation[conversation_id]:
                duplicates += 1
                continue
            analysis.execute(
                insert_sql,
                (
                    row["conversation_id"], row["message_timestamp"], row["message_datetime"],
                    row["he_number"], row["customer_number"], row["sender_number"],
                    row["sender_type"], row["message_type"], row["message_content"],
                    row["message_content_lower"], row["message_content_normalized"], row["source_file"],
                ),
            )
            existing_by_conversation[conversation_id].add(key)
            inserted += 1
        analysis.commit()

        with _connection() as conn:
            conn.execute(
                """
                UPDATE import_jobs SET status = 'indexing', inserted_messages = ?,
                    duplicate_messages = ?, updated_at = ? WHERE job_id = ?
                """,
                (inserted, duplicates, _now(), job_id),
            )
        if inserted:
            refresh_analysis_artifacts_preserving_enrichment(analysis)

    with _connection() as conn:
        failed = conn.execute(
            "SELECT failed_pairs FROM import_jobs WHERE job_id = ?", (job_id,)
        ).fetchone()[0]
        final_status = "completed_with_errors" if failed else "completed"
        conn.execute(
            """
            UPDATE import_jobs SET status = ?, inserted_messages = ?, duplicate_messages = ?,
                finished_at = ?, updated_at = ? WHERE job_id = ?
            """,
            (final_status, inserted, duplicates, _now(), _now(), job_id),
        )


initialize_import_store()
